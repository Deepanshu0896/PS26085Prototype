"""
Create AGASTYA_FINAL_TECHNICAL_REPORT_SIH2026.docx
Smart India Hackathon 2026 - Problem Statement 26085
Ministry of Earth Sciences (MoES)
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

WORKSPACE = Path(r"c:\Users\Deepanshu\OneDrive\Desktop\prototype26085")
DOCX_OUT = WORKSPACE / "AGASTYA_FINAL_TECHNICAL_REPORT_SIH2026.docx"

doc = docx.Document()

# Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

# Colors
COLOR_PRIMARY = RGBColor(0, 59, 92)     # #003B5C
COLOR_SECONDARY = RGBColor(2, 132, 199) # #0284C7
COLOR_TEXT = RGBColor(30, 41, 59)       # #1E293B
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B
COLOR_GREEN = RGBColor(21, 128, 61)     # #15803D
COLOR_ALERT = RGBColor(220, 38, 38)     # #DC2626

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(text, is_alert=False):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_background(c, "FEF2F2" if is_alert else "F0F9FF")
    set_cell_margins(c, 100, 100, 150, 150)
    p = c.paragraphs[0]
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    r.font.color.rgb = COLOR_ALERT if is_alert else COLOR_PRIMARY
    doc.add_paragraph()

# Header
p_title = doc.add_paragraph()
r_t = p_title.add_run("AGASTYA: Comprehensive Final Technical Report\n")
r_t.font.name = "Arial"
r_t.font.size = Pt(20)
r_t.font.bold = True
r_t.font.color.rgb = COLOR_PRIMARY

r_sub = p_title.add_run("Physics-Guided Urban Drainage Coupling, PySewer Integration & Constrained Emergency Routing\n")
r_sub.font.size = Pt(11)
r_sub.font.bold = True
r_sub.font.color.rgb = COLOR_SECONDARY

r_meta = p_title.add_run("Smart India Hackathon 2026 · Problem Statement 26085 · Ministry of Earth Sciences (MoES)")
r_meta.font.size = Pt(9.5)
r_meta.font.color.rgb = COLOR_MUTED

doc.add_paragraph()

# Executive Summary
h1 = doc.add_heading(level=1)
h1.add_run("1. Executive Summary & Problem Formulation").font.color.rgb = COLOR_PRIMARY

p_ex = doc.add_paragraph()
p_ex.add_run(
    "Urban flood nowcasting in India faces an operational disconnect: the India Meteorological Department (IMD) delivers "
    "macro-scale rainfall predictions (e.g., '60 mm across Central Delhi'), yet municipal authorities and ambulance dispatchers "
    "cannot determine which underpass or intersection will submerge. AGASTYA (Adaptive Geospatial Analysis & Stormwater Topology "
    "Yield Architecture) bridges this gap by coupling real-time Open-Meteo precipitation with a high-resolution digital terrain "
    "slope (CartoDEM 10m) and synthetic underground drainage conduits generated via PySewer. "
    "It delivers sub-second inundation predictions (<4.5ms), interactive multi-node choke modeling, storm duration accumulation scaling, "
    "and dynamic Dijkstra safe ambulance routing calculated instantly (<2ms) enforcing a 15.0 cm vehicular water clearance threshold."
)

# Table 1: Data segregation
h2 = doc.add_heading(level=2)
h2.add_run("Empirical vs Synthetic Data Segregation").font.color.rgb = COLOR_SECONDARY

tbl_data = doc.add_table(rows=5, cols=4)
tbl_data.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Data Layer", "Data Source & Resolution", "Classification", "Functional Role in AGASTYA"]
for idx, h in enumerate(headers):
    c = tbl_data.rows[0].cells[idx]
    set_cell_background(c, "003B5C")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

data_rows = [
    ("Precipitation Input", "Open-Meteo API / IMD Radar (0.25°)", "Empirical Real", "Supplies real-time hourly rainfall intensity (0-100 mm/hr) and storm duration curves."),
    ("Topography (DEM)", "ISRO CartoDEM 10m / SRTM 30m", "Empirical Real", "Defines slope and watershed boundaries; reveals Minto Bridge underpass sag (210.5m) vs CP ridge (216.5m)."),
    ("Road Network", "OpenStreetMap / OSMnx graph extraction", "Empirical Real", "Supplies road segment lengths, coordinates, junction connectivity, and emergency traversal corridors."),
    ("Drainage Conduits", "PySewer Synthetic Topology (JOSS 2024)", "Synthetic Proxy", "Synthesizes CPHEEO-compliant gravity pipes (300-1800mm, n=0.015) solving India's missing pipe data problem.")
]

for row_idx, row in enumerate(data_rows, start=1):
    r_elem = tbl_data.rows[row_idx]
    for col_idx, val in enumerate(row):
        c = r_elem.cells[col_idx]
        set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(c, 60, 60, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
        elif col_idx == 2:
            r.font.bold = True
            r.font.color.rgb = COLOR_GREEN if "Real" in val else COLOR_SECONDARY

doc.add_paragraph()

# Mathematical formulation
h2_math = doc.add_heading(level=1)
h2_math.add_run("2. Mathematical Formulation of Hydraulic Engine").font.color.rgb = COLOR_PRIMARY

p_math = doc.add_paragraph()
p_math.add_run(
    "1. Rational Overland Inflow: Q_in,i = (C_i · I · A_i) / 360 [m³/s], where C_i = 0.85, I = rainfall rate (mm/hr), A_i = catchment area.\n"
    "2. Manning Pipe Conveyance Capacity: Q_cap,ij = (1 / n) · A_pipe · R^(2/3) · S^(1/2) [m³/s], where n = 0.015, R = D/4, S = hydraulic slope.\n"
    "3. Overland Inundation Surcharge: Q_surcharge,i = max(0, Q_in,i - Σ Q_cap,out + Q_backwater_in).\n"
    "4. Depth Ponding: Water_Depth_i (cm) = (Q_surcharge,i · Δt · Duration_Scale / Ponding_Area_i) · 100.\n"
    "• Validation at Minto Sag (210.5m): 75 mm/hr downpour accumulates 38.3 cm (CRITICAL); 0 mm/hr yields strictly 0.0 cm (SAFE)."
)
p_math.runs[0].font.name = "Consolas"
p_math.runs[0].font.size = Pt(8.5)

# Routing & Safety Gates
h3_route = doc.add_heading(level=1)
h3_route.add_run("3. Constrained Dijkstra Emergency Routing & Safety Lifecycle").font.color.rgb = COLOR_PRIMARY

add_callout(
    "LIFE-SAFETY CLEARANCE THRESHOLD (15.0 cm): Emergency ambulances (Force Traveller, Tata Winger) operate with engine intake "
    "clearances of 200-250 mm. Traversal through water depths >15.0 cm causes mechanical hydrolock and engine failure. "
    "AGASTYA enforces strict pre-Dijkstra validation: any dispatch request with origin or destination >15.0 cm halts with "
    "structured ORIGIN_UNSAFE or DESTINATION_UNSAFE warnings, preventing vehicles from entering flooded water.",
    is_alert=True
)

# 53-Test Suite
h4_test = doc.add_heading(level=1)
h4_test.add_run("4. 53-Test Automated Regression Suite").font.color.rgb = COLOR_PRIMARY

tbl_tests = doc.add_table(rows=6, cols=4)
tbl_tests.alignment = WD_TABLE_ALIGNMENT.CENTER
t_headers = ["Subsystem", "Tests", "Key Verification Assertions", "Result"]
for idx, h in enumerate(t_headers):
    c = tbl_tests.rows[0].cells[idx]
    set_cell_background(c, "0284C7")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

test_rows = [
    ("Baseline Hydraulic Physics", "12", "Zero rain yields strictly 0.0 cm depth; Manning capacity monotonically rises with pipe diameter; Rational runoff scales linearly.", "12/12 PASS"),
    ("Sag Basin Concentration", "10", "Lowest elevation (210.5m) accumulates maximum backwater surcharge (38.3 cm @ 75 mm/hr); higher elevations (216.5m) drain safely (<2 cm).", "10/10 PASS"),
    ("Multi-Node Choke Engine", "11", "Simultaneous manhole blockage; upstream backwater propagation; unblock state restoration; zero side-effects when choke switch is off.", "11/11 PASS"),
    ("Endpoint Safety & Edge Pruning", "12", "Submerged origin returns ORIGIN_UNSAFE; submerged destination returns DESTINATION_UNSAFE; flooded edges (>15 cm) pruned from graph.", "12/12 PASS"),
    ("Offline Client Parity & Schema", "8", "Full client-side TypeScript physical and Dijkstra solver mirrors Python backend; Pydantic v2 schemas pass serialization.", "8/8 PASS")
]

for row_idx, row in enumerate(test_rows, start=1):
    r_elem = tbl_tests.rows[row_idx]
    for col_idx, val in enumerate(row):
        c = r_elem.cells[col_idx]
        set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(c, 60, 60, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
        elif col_idx == 3:
            r.font.bold = True
            r.font.color.rgb = COLOR_GREEN

doc.add_paragraph()

# Visual Evidence Screenshots embedded into DOCX!
h5_vis = doc.add_heading(level=1)
h5_vis.add_run("5. Live Prototype UI Verification Evidence").font.color.rgb = COLOR_PRIMARY

img_a = WORKSPACE / "test_a_flooded_origin.png"
img_b = WORKSPACE / "test_b_flooded_destination.png"
img_c = WORKSPACE / "test_c_rerouted_corridor.png"

for img_path, caption in [
    (img_a, "Figure 1: Test Case A — Submerged Origin Endpoint (Depth = 17.5 cm > 15 cm threshold). System halts dispatch with ORIGIN_UNSAFE alert."),
    (img_b, "Figure 2: Test Case B — Submerged Destination Endpoint (Depth = 19.8 cm > 15 cm threshold). System triggers DESTINATION_UNSAFE alert to reroute transport."),
    (img_c, "Figure 3: Test Case C — Flooded Sag Bypass (Minto Sag depth = 38.3 cm). AGASTYA prunes flooded underpass and computes safe green arterial bypass via DDU Marg East in <2ms.")
]:
    if img_path.exists():
        doc.add_picture(str(img_path), width=Inches(6.5))
        p_cap = doc.add_paragraph()
        r_c = p_cap.add_run(caption)
        r_c.font.size = Pt(8)
        r_c.font.italic = True
        r_c.font.color.rgb = COLOR_MUTED
        doc.add_paragraph()

# Market Comparison & Roadmaps
h6_comp = doc.add_heading(level=1)
h6_comp.add_run("6. Competitive Comparison vs Market Solutions").font.color.rgb = COLOR_PRIMARY

tbl_comp = doc.add_table(rows=9, cols=5)
tbl_comp.alignment = WD_TABLE_ALIGNMENT.CENTER
c_headers = ["Key Capability", "Traditional Flood Portals", "Commercial Flood APIs", "Municipal ICCC Dashboards", "AGASTYA (Our Solution)"]
for idx, h in enumerate(c_headers):
    c = tbl_comp.rows[0].cells[idx]
    set_cell_background(c, "003B5C")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

comp_data = [
    ("Street-Level Depth (cm)", "✕ (City-wide only)", "▲ (Satellite extent only)", "✕ (Qualitative text)", "✓ (Exact cm per node)"),
    ("Drainage Pipe Coupling", "✕ (Ignored)", "✕ (Overland only)", "✕ (Static asset list)", "✓ (Manning 1D physics)"),
    ("Missing Pipe Synthesis", "✕ (Requires CAD)", "✕ (Requires sensors)", "✕ (Missing data barrier)", "✓ (PySewer AI synthesis)"),
    ("Safe Ambulance Routing", "✕ (No routing)", "✕ (Historical maps)", "▲ (Manual radio dispatch)", "✓ (Dijkstra <15cm gate)"),
    ("Submerged Endpoint Refusal", "✕ (No safety checks)", "✕ (No vehicle logic)", "✕ (Unsafe dispatch)", "✓ (ORIGIN / DEST UNSAFE)"),
    ("Multi-Node Choke Simulation", "✕ (Static models)", "✕ (No clog modeling)", "✕ (Reactive complaints)", "✓ (Interactive toggle)"),
    ("100% Offline Resilience", "✕ (Crashes offline)", "✕ (Cloud-dependent)", "✕ (Server-dependent)", "✓ (Client-side engine)"),
    ("Deployment Cost", "High (₹1–2 Cr consulting)", "High (Recurring SaaS $)", "Extreme (₹5–20 Cr SCADA)", "✓ (<₹10 Lakhs/city)")
]

for row_idx, row in enumerate(comp_data, start=1):
    r_elem = tbl_comp.rows[row_idx]
    for col_idx, val in enumerate(row):
        c = r_elem.cells[col_idx]
        set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(c, 60, 60, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY
        elif col_idx == 4:
            r.font.bold = True
            r.font.color.rgb = COLOR_GREEN

doc.add_paragraph()

# 36-hr roadmap & team contribution
h7_team = doc.add_heading(level=1)
h7_team.add_run("7. 6-Member Team Contribution & 36-Hour Hackathon Roadmap").font.color.rgb = COLOR_PRIMARY

tbl_team = doc.add_table(rows=7, cols=4)
tbl_team.alignment = WD_TABLE_ALIGNMENT.CENTER
team_heads = ["Member Name", "Assigned Role", "Core Responsibilities", "Delivered Artifacts"]
for idx, h in enumerate(team_heads):
    c = tbl_team.rows[0].cells[idx]
    set_cell_background(c, "0284C7")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    r = p.add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(255, 255, 255)

team_members = [
    ("Deepanshu (Lead)", "Lead System Architect", "End-to-end architecture, FastAPI backend, pitch delivery, technical defense.", "Core platform, technical report, pitch script."),
    ("Member 2", "Hydrologic / GIS Engineer", "CartoDEM slope sampling, PySewer pipe synthesis, Manning capacity modeling.", "25-node topology graph, elevation model."),
    ("Member 3", "Frontend UI/UX Specialist", "React 19 dashboard, Leaflet mapping visualizer, color tokens, interactive controls.", "Interactive operator UI, glassmorphic HUD."),
    ("Member 4", "Algorithm & Routing Lead", "Dijkstra emergency solver, vehicle water clearance gates, edge pruning.", "Routing engine, endpoint safety logic."),
    ("Member 5", "QA & Reliability Engineer", "Automated pytest suite (53 tests), regression auditing, browser verification.", "Passing test suite, regression matrix, demo videos."),
    ("Member 6", "Policy & Research Analyst", "Economic loss data collection, NDMA/CPHEEO compliance, market benchmarking.", "Facts sheet, comparison matrix, references.")
]

for row_idx, row in enumerate(team_members, start=1):
    r_elem = tbl_team.rows[row_idx]
    for col_idx, val in enumerate(row):
        c = r_elem.cells[col_idx]
        set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(c, 60, 60, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(8)
        if col_idx == 0:
            r.font.bold = True
            r.font.color.rgb = COLOR_PRIMARY

doc.save(str(DOCX_OUT))
print(f"Successfully generated DOCX: {DOCX_OUT} (Size: {DOCX_OUT.stat().st_size} bytes)")
