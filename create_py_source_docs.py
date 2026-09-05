"""
Create PDF and DOCX documentation of the Python source code files:
1. compile_final_technical_report_pdf.py
2. create_presentation_docx.py
Produces:
- AGASTYA_PYTHON_SOURCE_CODE_SIH2026.docx
- AGASTYA_PYTHON_SOURCE_CODE_SIH2026.pdf
"""

import os
import subprocess
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

WORKSPACE = Path(r"c:\Users\Deepanshu\OneDrive\Desktop\prototype26085")
DOCX_OUT = WORKSPACE / "AGASTYA_PYTHON_SOURCE_CODE_SIH2026.docx"
PDF_OUT = WORKSPACE / "AGASTYA_PYTHON_SOURCE_CODE_SIH2026.pdf"

doc = docx.Document()
for section in doc.sections:
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

COLOR_PRIMARY = RGBColor(0, 59, 92)     # #003B5C
COLOR_SECONDARY = RGBColor(2, 132, 199) # #0284C7
COLOR_TEXT = RGBColor(30, 41, 59)       # #1E293B

p = doc.add_paragraph()
r_t = p.add_run("AGASTYA: Python Source Code Documentation\n")
r_t.font.name = "Arial"
r_t.font.size = Pt(18)
r_t.font.bold = True
r_t.font.color.rgb = COLOR_PRIMARY

r_s = p.add_run("Automated Compilers for Final Technical Report PDF and SIH 2026 Presentation Content\n")
r_s.font.size = Pt(10.5)
r_s.font.bold = True
r_s.font.color.rgb = COLOR_SECONDARY

r_m = p.add_run("Smart India Hackathon 2026 · Problem Statement 26085 · Ministry of Earth Sciences (MoES)")
r_m.font.size = Pt(9)

doc.add_paragraph()

files_to_include = [
    ("create_presentation_docx.py", "Presentation Content & Slide Blueprint Generator (.docx)"),
    ("compile_final_technical_report_pdf.py", "Comprehensive Final Technical Report Compiler (.pdf)")
]

for filename, desc in files_to_include:
    fpath = WORKSPACE / filename
    if not fpath.exists():
        continue
    
    h = doc.add_heading(level=1)
    r_h = h.add_run(f"Source Code: {filename}")
    r_h.font.color.rgb = COLOR_PRIMARY
    
    p_desc = doc.add_paragraph()
    r_d = p_desc.add_run(f"Purpose: {desc}\nFile Path: {fpath}")
    r_d.font.size = Pt(8.5)
    r_d.font.italic = True
    
    with open(fpath, "r", encoding="utf-8") as f:
        code_text = f.read()
    
    # Add code block
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
    tcPr.append(shd)
    
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="120" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/><w:left w:w="160" w:type="dxa"/><w:right w:w="160" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)
    
    p_code = cell.paragraphs[0]
    p_code.paragraph_format.line_spacing = 1.05
    r_code = p_code.add_run(code_text)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(6.8)
    r_code.font.color.rgb = RGBColor(56, 189, 248) # #38bdf8 light cyan
    
    doc.add_paragraph()

doc.save(str(DOCX_OUT))
print(f"Successfully generated DOCX: {DOCX_OUT} (Size: {DOCX_OUT.stat().st_size} bytes)")

# Convert to PDF via Word COM
import win32com.client
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    wdoc = word.Documents.Open(str(DOCX_OUT))
    wdoc.SaveAs(str(PDF_OUT), FileFormat=17)
    wdoc.Close()
    print(f"Successfully generated PDF: {PDF_OUT} (Size: {PDF_OUT.stat().st_size} bytes)")
finally:
    word.Quit()
